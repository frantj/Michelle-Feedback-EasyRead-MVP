#!/usr/bin/env python3
"""Render an Easy Read JSON file to a Word document, with an image beside each sentence.

Usage:
  python3 build_document.py input.json output.docx [--pdf] [--font Arial] [--size 14] [--image-cm 3]

The layout follows easyreadgenerator.com: title, summary box, then each section
heading followed by rows of image (left) and sentence (right). A final "Image
credits" line lists the sources of the images used.

--pdf also writes a PDF next to the .docx, using LibreOffice (soffice).
"""

import argparse
import hashlib
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from easy_read_common import IMAGES_DIR, load_image_map, load_json, resolve_keyword, schema_errors

try:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.image.image import Image as DocxImage
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    sys.exit("ERROR: python-docx is not installed. Run: pip install python-docx")

AMBER = "F7A823"
AMBER_LIGHT = "FFF6E0"
GRAY_BORDER = "E0E0E0"
TEXT_WIDTH = Cm(17)  # A4 (21 cm) minus 2 cm margins
RASTER_TYPES ={".png", ".jpg", ".jpeg"}
PNG_CACHE = Path(tempfile.gettempdir()) / "easy-read-png-cache"


# --- Images ---

def image_to_raster(path):
    """Return a PNG/JPEG path for the image, converting SVG with cairosvg (cached)."""
    if path.suffix.lower() in RASTER_TYPES:
        return path
    try:
        import cairosvg
    except (ImportError, OSError) as e:
        sys.exit("ERROR: cairosvg is not available, so SVG images cannot be converted. "
                 f"Run: pip install cairosvg (it also needs the Cairo library). Details: {e}")
    PNG_CACHE.mkdir(exist_ok=True)
    digest = hashlib.sha1(path.read_bytes()).hexdigest()[:16]
    out = PNG_CACHE / f"{path.stem}-{digest}.png"
    if not out.exists():
        cairosvg.svg2png(url=str(path), write_to=str(out), output_width=360)
    return out


def fit_size(path, box_cm):
    """Width and height that fit the image inside a box_cm square, keeping its shape."""
    img = DocxImage.from_file(str(path))
    w, h = img.px_width, img.px_height
    if w >= h:
        return Cm(box_cm), Cm(box_cm * h / w)
    return Cm(box_cm * w / h), Cm(box_cm)


# --- Word formatting helpers ---

# Word expects properties in schema order; appended elements are moved into place
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
             "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
             "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
             "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
             "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
             "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
              "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"]


def sort_children(el, order):
    rank = {qn(f"w:{name}"): i for i, name in enumerate(order)}
    children = sorted(el, key=lambda c: rank.get(c.tag, len(order)))
    for c in children:
        el.remove(c)
        el.append(c)


def set_style_font(style, font, size_pt, bold=None):
    style.font.size = Pt(size_pt)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.italic = False
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    # Theme font attributes override named fonts, so remove them
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{attr}"), None)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), font)


def paragraph_border(paragraph, side, color, size_eighths, space=4):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size_eighths))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pbdr.append(el)
    sort_children(ppr, PPR_ORDER)


def cell_shading(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr = cell._tc.get_or_add_tcPr()
    tcpr.append(shd)
    sort_children(tcpr, TCPR_ORDER)


def cell_border(cell, side, color, size_eighths):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size_eighths))
    el.set(qn("w:color"), color)
    borders.append(el)
    sort_children(tcpr, TCPR_ORDER)


def fix_table_widths(table, widths):
    """Set fixed column widths so every app (Word, Google Docs, LibreOffice) uses them."""
    table.autofit = False
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(sum(w.twips for w in widths)))
    tblw.set(qn("w:type"), "dxa")  # autofit = False already set the fixed layout
    for col, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(w.twips))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w


def keep_row_together(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def set_alt_text(inline_shape, text):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", text)
    doc_pr.set("title", text)


# --- Document ---

def setup_document(font, size):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)  # A4
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2))

    normal = doc.styles["Normal"]
    set_style_font(normal, font, size)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for name, pts in (("Heading 1", size * 2), ("Heading 2", size * 1.45)):
        style = doc.styles[name]
        set_style_font(style, font, pts, bold=True)
        style.paragraph_format.space_before = Pt(18)
        style.paragraph_format.space_after = Pt(10)
        style.paragraph_format.keep_with_next = True
    return doc


def add_summary(doc, summary, size):
    """A shaded box with an amber left edge, like the site's summary box."""
    table = doc.add_table(rows=1, cols=1)
    fix_table_widths(table, [TEXT_WIDTH])
    cell = table.cell(0, 0)
    cell_shading(cell, AMBER_LIGHT)
    cell_border(cell, "left", AMBER, 36)
    label = cell.paragraphs[0]
    run = label.add_run("Summary")
    run.bold = True
    run.font.size = Pt(size * 1.15)
    label.paragraph_format.space_before = Pt(8)
    label.paragraph_format.space_after = Pt(2)
    body = cell.add_paragraph(summary)
    body.paragraph_format.space_after = Pt(10)
    doc.add_paragraph()


def add_section(doc, section, image_map, args, stats):
    heading = doc.add_heading(section["heading"], level=2)
    paragraph_border(heading, "bottom", AMBER, 18)

    sentences = section["sentences"]
    table = doc.add_table(rows=len(sentences), cols=2)
    image_col = Cm(args.image_cm + 0.8)
    fix_table_widths(table, [image_col, Cm(TEXT_WIDTH.cm - image_col.cm)])

    for i, (row, sentence) in enumerate(zip(table.rows, sentences)):
        keep_row_together(row)
        img_cell, text_cell = row.cells
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i < len(sentences) - 1:
                cell_border(cell, "bottom", GRAY_BORDER, 4)

        text_para = text_cell.paragraphs[0]
        text_para.text = sentence["text"]
        text_para.paragraph_format.space_before = Pt(6)

        stats["sentences"] += 1
        kw = sentence.get("imageKeyword") or ""
        match, how = resolve_keyword(kw, image_map)
        if how == "none":
            stats["missing"].append(f'"{kw}" – {sentence["text"]}' if kw else f'(no keyword) – {sentence["text"]}')
            continue
        if how == "fallback":
            stats["fallbacks"].append(f'"{kw}" -> "{match}" – {sentence["text"]}')

        entry = image_map[match]
        src = IMAGES_DIR / entry["file"]
        if not src.is_file():
            stats["missing"].append(f'"{kw}" (file {entry["file"]} not found) – {sentence["text"]}')
            continue
        raster = image_to_raster(src)
        width, height = fit_size(raster, args.image_cm)
        img_para = img_cell.paragraphs[0]
        img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_para.paragraph_format.line_spacing = 1.0
        img_para.paragraph_format.space_before = Pt(6)
        shape = img_para.add_run().add_picture(str(raster), width=width, height=height)
        set_alt_text(shape, entry.get("alt") or match)
        stats["placed"] += 1
        stats["sources"].add(entry.get("source"))
        stats["svg_converted"] |= src.suffix.lower() == ".svg"

    doc.add_paragraph()


def add_credits(doc, attribution, stats, size):
    sources = [attribution[s] for s in sorted(s for s in stats["sources"] if s in attribution)]
    if not sources:
        return
    parts = [f'{s["name"]} ({s["license"]}, {s["url"]})' for s in sources]
    text = "Image credits: pictures from " + "; ".join(parts) + "."
    if stats["svg_converted"]:
        text += " Some pictures were converted from SVG to PNG."
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(max(10, size - 3))
    p.paragraph_format.space_before = Pt(18)
    paragraph_border(p, "top", GRAY_BORDER, 6)


def convert_to_pdf(docx_path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
        soffice = str(mac) if mac.exists() else None
    if not soffice:
        print("WARNING: LibreOffice (soffice) not found, so no PDF was made. The .docx is ready.")
        return None
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir",
                    str(docx_path.parent), str(docx_path)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    pdf = docx_path.with_suffix(".pdf")
    return pdf if pdf.exists() else None


def main():
    parser = argparse.ArgumentParser(description="Render Easy Read JSON to a Word document.")
    parser.add_argument("input")
    parser.add_argument("output")
    parser.add_argument("--pdf", action="store_true", help="also write a PDF (needs LibreOffice)")
    parser.add_argument("--font", default="Arial")
    parser.add_argument("--size", type=float, default=14, help="body text size in points (default 14)")
    parser.add_argument("--image-cm", type=float, default=3, help="image size in cm (default 3)")
    args = parser.parse_args()

    doc_json = load_json(args.input)
    errors = schema_errors(doc_json)
    if errors:
        sys.exit("ERROR: the JSON does not match the Easy Read format:\n  " + "\n  ".join(errors))

    image_map, attribution = load_image_map()
    out = Path(args.output)
    if out.suffix.lower() != ".docx":
        out = out.with_suffix(".docx")
    out.parent.mkdir(parents=True, exist_ok=True)

    stats = {"sentences": 0, "placed": 0, "fallbacks": [], "missing": [],
             "sources": set(), "svg_converted": False}
    doc = setup_document(args.font, args.size)
    title = doc.add_heading(doc_json["title"], level=1)
    paragraph_border(title, "bottom", AMBER, 24, space=6)
    add_summary(doc, doc_json["summary"], args.size)
    for section in doc_json["sections"]:
        add_section(doc, section, image_map, args, stats)
    add_credits(doc, attribution, stats, args.size)
    doc.core_properties.title = doc_json["title"]
    doc.save(out)

    print(f"Wrote {out}")
    if args.pdf:
        pdf = convert_to_pdf(out)
        if pdf:
            print(f"Wrote {pdf}")
    print(f"Sentences: {stats['sentences']}  Images placed: {stats['placed']}  "
          f"Fallbacks: {len(stats['fallbacks'])}  Missing: {len(stats['missing'])}")
    for f in stats["fallbacks"]:
        print(f"  FALLBACK {f}")
    for m in stats["missing"]:
        print(f"  MISSING  {m}")
    if stats["fallbacks"] or stats["missing"]:
        print("Pick catalog keywords for these sentences and build again.")


if __name__ == "__main__":
    main()
